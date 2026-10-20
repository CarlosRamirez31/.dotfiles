#!/usr/bin/env python3
"""Local-first document retrieval with a normalized, auditable response."""

from __future__ import annotations

import argparse
import csv
import hashlib
import html
import json
import math
import os
import re
import sys
import unicodedata
import zipfile
from collections import Counter, defaultdict
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Iterable
from xml.etree import ElementTree as ET


SUPPORTED = {".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".html", ".htm", ".docx", ".pptx", ".pdf"}
STOPWORDS = {
    "a", "al", "ante", "bajo", "como", "con", "cual", "cuál", "de", "del", "desde", "donde", "el", "en",
    "entre", "es", "esta", "este", "fue", "ha", "la", "las", "lo", "los", "más", "para", "pero", "por",
    "que", "qué", "se", "según", "sin", "sobre", "su", "sus", "un", "una", "y", "the", "a", "an", "of",
    "to", "in", "is", "are", "was", "were", "for", "and", "or", "on", "with", "what", "which", "who",
}
WORD_RE = re.compile(r"[A-Za-zÁÉÍÓÚÜÑáéíóúüñ0-9][A-Za-zÁÉÍÓÚÜÑáéíóúüñ0-9._-]*")


def hermes_home() -> Path:
    explicit = os.environ.get("HERMES_HOME")
    if explicit:
        return Path(explicit).expanduser().resolve()
    if os.name == "nt":
        return Path(os.environ.get("LOCALAPPDATA", str(Path.home() / "AppData/Local"))) / "hermes"
    return Path.home() / ".hermes"


def normalize(value: str) -> str:
    decomposed = unicodedata.normalize("NFKD", value.casefold())
    return "".join(char for char in decomposed if not unicodedata.combining(char))


def tokens(value: str) -> list[str]:
    return [token for token in (normalize(match.group(0)) for match in WORD_RE.finditer(value)) if token not in STOPWORDS and len(token) > 1]


class _HTMLText(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        if data.strip():
            self.parts.append(data.strip())


def _clean(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def _text_chunks(text: str, page: int | None = None, size: int = 1400, overlap: int = 180) -> list[dict[str, Any]]:
    paragraphs = [_clean(part) for part in re.split(r"\n\s*\n|(?<=[.!?])\s+(?=[A-ZÁÉÍÓÚÜÑ0-9])", text) if _clean(part)]
    chunks: list[dict[str, Any]] = []
    current = ""
    for paragraph in paragraphs:
        if current and len(current) + len(paragraph) + 1 > size:
            chunks.append({"text": current, "page": page})
            current = f"{current[-overlap:]} {paragraph}" if overlap else paragraph
        else:
            current = f"{current} {paragraph}".strip()
    if current:
        chunks.append({"text": current, "page": page})
    return chunks


def _read_plain(path: Path) -> list[dict[str, Any]]:
    return _text_chunks(path.read_text(encoding="utf-8-sig", errors="replace"))


def _read_csv(path: Path) -> list[dict[str, Any]]:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    rows: list[str] = []
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        for row in csv.reader(handle, delimiter=delimiter):
            rows.append(" | ".join(cell.strip() for cell in row))
    return _text_chunks("\n".join(rows))


def _read_json(path: Path) -> list[dict[str, Any]]:
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    return _text_chunks(json.dumps(data, ensure_ascii=False, indent=2))


def _read_html(path: Path) -> list[dict[str, Any]]:
    parser = _HTMLText()
    parser.feed(path.read_text(encoding="utf-8-sig", errors="replace"))
    return _text_chunks(html.unescape("\n".join(parser.parts)))


def _read_docx(path: Path) -> list[dict[str, Any]]:
    from docx import Document

    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return _text_chunks("\n\n".join(parts))


def _read_pptx(path: Path) -> list[dict[str, Any]]:
    chunks: list[dict[str, Any]] = []
    with zipfile.ZipFile(path) as archive:
        slides = sorted(name for name in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", name))
        for index, name in enumerate(slides, 1):
            root = ET.fromstring(archive.read(name))
            text = " ".join(node.text or "" for node in root.iter() if node.tag.endswith("}t"))
            chunks.extend(_text_chunks(text, page=index))
    return chunks


def _read_pdf(path: Path) -> list[dict[str, Any]]:
    from pypdf import PdfReader

    chunks: list[dict[str, Any]] = []
    for page_number, page in enumerate(PdfReader(str(path)).pages, 1):
        text = page.extract_text() or ""
        chunks.extend(_text_chunks(text, page=page_number))
    return chunks


READERS = {
    ".txt": _read_plain, ".md": _read_plain, ".markdown": _read_plain,
    ".csv": _read_csv, ".tsv": _read_csv, ".json": _read_json,
    ".html": _read_html, ".htm": _read_html, ".docx": _read_docx,
    ".pptx": _read_pptx, ".pdf": _read_pdf,
}


def source_signature(files: Iterable[Path], root: Path) -> str:
    digest = hashlib.sha256()
    for path in files:
        stat = path.stat()
        digest.update(str(path.relative_to(root)).encode("utf-8"))
        digest.update(f"{stat.st_size}:{stat.st_mtime_ns}".encode("ascii"))
    return digest.hexdigest()


def build_index(sources_dir: Path, cache_path: Path) -> dict[str, Any]:
    sources_dir.mkdir(parents=True, exist_ok=True)
    files = sorted(path for path in sources_dir.rglob("*") if path.is_file() and path.suffix.lower() in SUPPORTED)
    chunks: list[dict[str, Any]] = []
    errors: list[dict[str, str]] = []
    for path in files:
        try:
            for chunk_number, chunk in enumerate(READERS[path.suffix.lower()](path), 1):
                text = chunk["text"]
                if not text.strip():
                    continue
                token_counts = Counter(tokens(text))
                source = str(path.relative_to(sources_dir))
                source_id = "LOCAL-" + hashlib.sha1(source.encode("utf-8")).hexdigest()[:8].upper()
                chunks.append({
                    "text": text,
                    "source": source,
                    "path": str(path.resolve()),
                    "page": chunk.get("page"),
                    "chunk": chunk_number,
                    "source_id": source_id,
                    "tokens": dict(token_counts),
                    "length": sum(token_counts.values()),
                })
        except Exception as exc:
            errors.append({"source": str(path), "error": str(exc)})
    document_frequency: dict[str, int] = defaultdict(int)
    for chunk in chunks:
        for token in chunk["tokens"]:
            document_frequency[token] += 1
    index = {
        "version": 1,
        "sources_dir": str(sources_dir.resolve()),
        "signature": source_signature(files, sources_dir),
        "file_count": len(files),
        "chunk_count": len(chunks),
        "chunks": chunks,
        "document_frequency": dict(document_frequency),
        "errors": errors,
    }
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    cache_path.write_text(json.dumps(index, ensure_ascii=False, indent=2), encoding="utf-8")
    return index


def load_index(sources_dir: Path, cache_path: Path) -> dict[str, Any]:
    files = sorted(path for path in sources_dir.rglob("*") if path.is_file() and path.suffix.lower() in SUPPORTED) if sources_dir.exists() else []
    if cache_path.exists():
        try:
            cached = json.loads(cache_path.read_text(encoding="utf-8"))
            if cached.get("signature") == source_signature(files, sources_dir):
                return cached
        except (OSError, ValueError):
            pass
    return build_index(sources_dir, cache_path)


def search_local_documents(query: str, top_k: int = 6, sources_dir: Path | None = None, cache_path: Path | None = None) -> dict[str, Any]:
    home = hermes_home()
    sources = sources_dir or home / "academic" / "sources"
    cache = cache_path or home / "academic" / "cache" / "local_index.json"
    index = load_index(sources, cache)
    query_tokens = tokens(query)
    query_counts = Counter(query_tokens)
    total_chunks = max(1, int(index.get("chunk_count", 0)))
    average_length = sum(chunk.get("length", 0) for chunk in index.get("chunks", [])) / total_chunks or 1.0
    scored: list[dict[str, Any]] = []
    for chunk in index.get("chunks", []):
        score = 0.0
        length = max(1, chunk.get("length", 0))
        for token, query_frequency in query_counts.items():
            tf = chunk.get("tokens", {}).get(token, 0)
            if not tf:
                continue
            df = index.get("document_frequency", {}).get(token, 0)
            idf = math.log(1 + (total_chunks - df + 0.5) / (df + 0.5))
            k1, b = 1.5, 0.75
            bm25 = idf * (tf * (k1 + 1)) / (tf + k1 * (1 - b + b * length / average_length))
            score += query_frequency * bm25
        normalized_text = normalize(chunk["text"])
        normalized_query = normalize(query).strip()
        if normalized_query and normalized_query in normalized_text:
            score += 2.0
        if score > 0:
            scored.append({
                "text": chunk["text"], "source": chunk["source"], "path": chunk["path"],
                "page": chunk.get("page"), "score": round(score, 6), "source_id": chunk["source_id"],
            })
    scored.sort(key=lambda item: (-item["score"], item["source"], item.get("page") or 0))
    result = scored[: max(1, top_k)]
    matched_tokens = {token for token in query_counts if any(token in chunk.get("tokens", {}) for chunk in index.get("chunks", []))}
    coverage = len(matched_tokens) / len(set(query_tokens)) if query_tokens else 0.0
    top_score = result[0]["score"] if result else 0.0
    sufficient = bool(result and coverage >= 0.6 and top_score >= 0.5)
    return {
        "query": query,
        "results": result,
        "score_direction": "higher_is_better",
        "assessment": {
            "local_sufficient": sufficient,
            "query_token_coverage": round(coverage, 3),
            "top_score": top_score,
            "known": [item["source_id"] for item in result] if sufficient else [],
            "missing": [] if sufficient else [query],
            "uncertain": [] if sufficient else ([item["source_id"] for item in result] if result else []),
        },
        "index": {key: index.get(key) for key in ("file_count", "chunk_count", "errors")},
    }


def main() -> int:
    home = hermes_home()
    parser = argparse.ArgumentParser()
    sub = parser.add_subparsers(dest="command", required=True)
    index_cmd = sub.add_parser("index")
    index_cmd.add_argument("--sources", type=Path, default=home / "academic" / "sources")
    index_cmd.add_argument("--cache", type=Path, default=home / "academic" / "cache" / "local_index.json")
    search_cmd = sub.add_parser("search")
    search_cmd.add_argument("query")
    search_cmd.add_argument("--top-k", type=int, default=6)
    search_cmd.add_argument("--sources", type=Path, default=home / "academic" / "sources")
    search_cmd.add_argument("--cache", type=Path, default=home / "academic" / "cache" / "local_index.json")
    search_cmd.add_argument("--json", action="store_true")
    health_cmd = sub.add_parser("health")
    health_cmd.add_argument("--sources", type=Path, default=home / "academic" / "sources")
    health_cmd.add_argument("--cache", type=Path, default=home / "academic" / "cache" / "local_index.json")
    args = parser.parse_args()
    if args.command == "index":
        result = build_index(args.sources, args.cache)
    elif args.command == "health":
        result = load_index(args.sources, args.cache)
        result = {"ok": not result.get("errors"), "file_count": result["file_count"], "chunk_count": result["chunk_count"], "errors": result["errors"]}
    else:
        result = search_local_documents(args.query, args.top_k, args.sources, args.cache)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if not (args.command == "health" and not result["ok"]) else 1


if __name__ == "__main__":
    sys.exit(main())
