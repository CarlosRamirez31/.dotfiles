#!/usr/bin/env python3
"""Create a reusable APA 7 Word template."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt


FONT = "Times New Roman"


def set_font(style, size: float = 12, bold: bool = False, italic: bool = False) -> None:
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    else:
        fonts.attrib.clear()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), FONT)
    color = style.element.get_or_add_rPr().find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        style.element.get_or_add_rPr().append(color)
    else:
        color.attrib.clear()
    color.set(qn("w:val"), "000000")
    for tag in ("w:spacing", "w:kern", "w:effect", "w:shadow", "w:outline"):
        inherited = rpr.find(qn(tag))
        if inherited is not None:
            rpr.remove(inherited)


def set_paragraph(style, *, before=0, after=0, first_line=0, left=0, alignment=None, keep_with_next=False) -> None:
    fmt = style.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    fmt.first_line_indent = Inches(first_line)
    fmt.left_indent = Inches(left)
    fmt.keep_with_next = keep_with_next
    if alignment is not None:
        fmt.alignment = alignment
    paragraph_properties = style.element.get_or_add_pPr()
    border = paragraph_properties.find(qn("w:pBdr"))
    if border is not None:
        paragraph_properties.remove(border)


def ensure_style(document: Document, name: str, base: str = "Normal"):
    try:
        return document.styles[name]
    except KeyError:
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles[base]
        return style


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def create(path: Path, sample: bool = False) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    # Perfil institucional UNICARIBE: 2.5 cm exactos en los cuatro lados.
    section.top_margin = section.bottom_margin = Cm(2.5)
    section.left_margin = section.right_margin = Cm(2.5)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)
    add_page_number(section.header.paragraphs[0])

    normal = document.styles["Normal"]
    set_font(normal)
    set_paragraph(normal, first_line=0.5)

    title = document.styles["Title"]
    set_font(title, bold=True)
    set_paragraph(title, alignment=WD_ALIGN_PARAGRAPH.CENTER, keep_with_next=True)

    subtitle = document.styles["Subtitle"]
    set_font(subtitle)
    set_paragraph(subtitle, alignment=WD_ALIGN_PARAGRAPH.CENTER)

    heading1 = document.styles["Heading 1"]
    set_font(heading1, bold=True)
    set_paragraph(heading1, before=0, after=0, alignment=WD_ALIGN_PARAGRAPH.CENTER, keep_with_next=True)

    heading2 = document.styles["Heading 2"]
    set_font(heading2, bold=True)
    set_paragraph(heading2, before=0, after=0, alignment=WD_ALIGN_PARAGRAPH.LEFT, keep_with_next=True)

    heading3 = document.styles["Heading 3"]
    set_font(heading3, bold=True, italic=True)
    set_paragraph(heading3, before=0, after=0, alignment=WD_ALIGN_PARAGRAPH.LEFT, keep_with_next=True)

    heading4 = document.styles["Heading 4"]
    set_font(heading4, bold=True)
    set_paragraph(heading4, before=0, after=0, first_line=0.5, alignment=WD_ALIGN_PARAGRAPH.LEFT, keep_with_next=True)

    heading5 = document.styles["Heading 5"]
    set_font(heading5, bold=True, italic=True)
    set_paragraph(heading5, before=0, after=0, first_line=0.5, alignment=WD_ALIGN_PARAGRAPH.LEFT, keep_with_next=True)

    block = ensure_style(document, "Block Quote")
    set_font(block)
    set_paragraph(block, first_line=0, left=0.5)

    references = ensure_style(document, "References")
    set_font(references)
    set_paragraph(references, first_line=-0.5, left=0.5)

    abstract = ensure_style(document, "Abstract")
    set_font(abstract)
    set_paragraph(abstract, first_line=0)

    keywords = ensure_style(document, "Keywords")
    set_font(keywords)
    set_paragraph(keywords, first_line=0.5)

    for name, bold, italic in (
        ("Table Number", True, False), ("Table Title", False, True),
        ("Figure Number", True, False), ("Figure Title", False, True),
        ("Note", False, False),
    ):
        style = ensure_style(document, name)
        set_font(style, bold=bold, italic=italic)
        set_paragraph(style, first_line=0, alignment=WD_ALIGN_PARAGRAPH.LEFT, keep_with_next=name != "Note")

    caption = document.styles["Caption"]
    set_font(caption, italic=True)
    set_paragraph(caption, first_line=0, alignment=WD_ALIGN_PARAGRAPH.LEFT, keep_with_next=True)

    if sample:
        for _ in range(5):
            document.add_paragraph("")
        document.add_paragraph("Título del trabajo académico", style="Title")
        for line in ("Nombre del estudiante", "Institución", "Curso", "Docente", "13 de octubre de 2026"):
            document.add_paragraph(line, style="Subtitle")
        document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        document.add_paragraph("Título del trabajo académico", style="Heading 1")
        document.add_paragraph("Este párrafo demuestra el estilo Normal, la doble separación y la sangría de primera línea configurada para el cuerpo del trabajo.")
        document.add_paragraph("Desarrollo", style="Heading 1")
        document.add_paragraph("Este es un ejemplo de contenido. La plantilla no debe utilizarse como fuente académica.")
        document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        document.add_paragraph("Referencias", style="Heading 1")
        document.add_paragraph("Apellido, A. A. (2024). Título de ejemplo. Editorial.", style="References")
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("output", type=Path)
    parser.add_argument("--sample", action="store_true")
    args = parser.parse_args()
    create(args.output, args.sample)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
