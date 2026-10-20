#!/usr/bin/env python3
"""Validate structural APA settings in a generated DOCX."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Cm, Inches, Pt


def close(value: int | None, expected: int, tolerance: int = 8000) -> bool:
    return value is not None and abs(value - expected) <= tolerance


def validate(path: Path) -> list[str]:
    document = Document(path)
    failures: list[str] = []
    if not document.paragraphs:
        failures.append("El documento no contiene párrafos")
        return failures
    required = {
        "Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "Heading 5",
        "Block Quote", "References", "Abstract", "Keywords", "Table Number", "Table Title",
        "Figure Number", "Figure Title", "Note", "Caption",
    }
    available = {style.name for style in document.styles}
    for style in sorted(required - available):
        failures.append(f"Falta el estilo {style}")
    for section_number, section in enumerate(document.sections, 1):
        portrait_letter = close(section.page_width, Inches(8.5)) and close(section.page_height, Inches(11))
        landscape_letter = close(section.page_width, Inches(11)) and close(section.page_height, Inches(8.5))
        if not (portrait_letter or landscape_letter):
            failures.append(f"Sección {section_number}: el papel no es Carta, en orientación vertical u horizontal")
        for name, value in (("superior", section.top_margin), ("inferior", section.bottom_margin), ("izquierdo", section.left_margin), ("derecho", section.right_margin)):
            if not close(value, Cm(2.5)):
                failures.append(f"Sección {section_number}: margen {name} no es 2.5 cm (perfil UNICARIBE)")
    normal = document.styles["Normal"]
    if normal.font.size and normal.font.size != Pt(12):
        failures.append("Normal no usa 12 pt")
    if normal.font.name and normal.font.name != "Times New Roman":
        failures.append("Normal no usa Times New Roman")
    if not close(normal.paragraph_format.first_line_indent, Inches(0.5)):
        failures.append("Normal no tiene sangría de primera línea de 1.27 cm")
    if normal.paragraph_format.line_spacing_rule != WD_LINE_SPACING.DOUBLE:
        failures.append("Normal no usa interlineado doble")
    if normal.paragraph_format.space_before not in (None, Pt(0)) or normal.paragraph_format.space_after not in (None, Pt(0)):
        failures.append("Normal agrega espacio antes o después del párrafo")
    if normal.paragraph_format.alignment not in (None, WD_ALIGN_PARAGRAPH.LEFT):
        failures.append("Normal no está alineado a la izquierda")
    references = document.styles["References"]
    if not close(references.paragraph_format.left_indent, Inches(0.5)) or not close(references.paragraph_format.first_line_indent, -Inches(0.5)):
        failures.append("References no tiene sangría francesa de 1.27 cm")
    reference_headings = [paragraph for paragraph in document.paragraphs if paragraph.text.strip().casefold() == "referencias"]
    for heading in reference_headings:
        if heading.alignment not in (WD_ALIGN_PARAGRAPH.CENTER, None) and heading.style.name == "Heading 1":
            failures.append("El título Referencias no está centrado")
    for section_number, section in enumerate(document.sections, 1):
        if " PAGE " not in section.header._element.xml:
            failures.append(f"Sección {section_number}: falta campo de número de página en el encabezado")
    return failures


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("document", type=Path)
    args = parser.parse_args()
    failures = validate(args.document)
    if failures:
        print("FAIL:")
        for failure in failures:
            print(f"- {failure}")
        return 1
    print("PASS")
    return 0


if __name__ == "__main__":
    sys.exit(main())
