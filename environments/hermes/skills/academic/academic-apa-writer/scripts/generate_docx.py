#!/usr/bin/env python3
"""Convert validated Markdown into a DOCX using the APA template."""

from __future__ import annotations

import argparse
import re
import unicodedata
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def sanitize_filename(title: str) -> str:
    value = unicodedata.normalize("NFKD", title)
    value = "".join(char for char in value if not unicodedata.combining(char))
    value = re.sub(r"[^A-Za-z0-9]+", "_", value).strip("_").casefold()
    return value[:100] or "trabajo_academico"


def clear_body(document: Document) -> None:
    body = document._element.body
    for element in list(body):
        if element.tag != qn("w:sectPr"):
            body.remove(element)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend([color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([run_properties, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text: str) -> None:
    pattern = re.compile(r"(!?\[[^\]]*\]\([^)]*\)|\*\*[^*]+\*\*|\*[^*]+\*)")
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        token = match.group(0)
        image = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", token)
        link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
        if image:
            image_path = Path(image.group(2)).expanduser()
            if image_path.exists():
                run = paragraph.add_run()
                run.add_picture(str(image_path), width=Inches(5.5))
            else:
                paragraph.add_run(f"[Imagen no disponible: {image.group(1)}]")
        elif link:
            add_hyperlink(paragraph, link.group(1), link.group(2))
        elif token.startswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        else:
            paragraph.add_run(token[1:-1]).italic = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def parse_table(lines: list[str]) -> list[list[str]]:
    return [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]


def set_cell_border(cell, *, top: bool = False, bottom: bool = False) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge_name, enabled in (("top", top), ("bottom", bottom)):
        if not enabled:
            continue
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "8")
        edge.set(qn("w:color"), "000000")


def add_table(document: Document, rows: list[list[str]]) -> None:
    if len(rows) > 1 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
        rows = [rows[0], *rows[2:]]
    width = max(len(row) for row in rows)
    table = document.add_table(rows=0, cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cells[index].paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            for run in paragraph.runs:
                run.bold = row_index == 0
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)
            set_cell_border(cells[index], top=row_index == 0, bottom=row_index in (0, len(rows) - 1))


def add_inline_heading(document: Document, text: str, level: int) -> None:
    paragraph = document.add_paragraph(style="Normal")
    heading, separator, continuation = text.partition(". ")
    heading_text = heading.rstrip(".") + "."
    run = paragraph.add_run(heading_text)
    run.bold = True
    run.italic = level == 5
    if separator and continuation:
        paragraph.add_run(" ")
        add_inline(paragraph, continuation)


def convert(markdown: Path, output: Path, template: Path, metadata: dict[str, str]) -> None:
    document = Document(template)
    clear_body(document)
    lines = markdown.read_text(encoding="utf-8-sig").splitlines()
    title = metadata.get("title") or next((line[2:].strip() for line in lines if line.startswith("# ")), markdown.stem)

    for _ in range(5):
        document.add_paragraph("")
    document.add_paragraph(title, style="Title")
    document.add_paragraph("")
    for key in ("author", "institution", "course", "instructor", "due_date"):
        if metadata.get(key):
            document.add_paragraph(metadata[key], style="Subtitle")
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    content_lines = lines[1:] if lines and lines[0].startswith("# ") else lines
    first_heading = next((line for line in content_lines if line.strip()), "")
    first_heading_text = re.sub(r"^#{1,6}\s+", "", first_heading).strip().casefold()
    starts_with_abstract = first_heading_text in {"resumen", "abstract"}
    in_references = False
    in_abstract = False
    body_has_content = False
    pending_visual_title: str | None = None
    if not starts_with_abstract:
        document.add_paragraph(title, style="Heading 1")
        body_has_content = True
    index = 0
    while index < len(content_lines):
        line = content_lines[index].rstrip()
        if line.startswith("|") and index + 1 < len(content_lines) and content_lines[index + 1].startswith("|"):
            table_lines = []
            while index < len(content_lines) and content_lines[index].startswith("|"):
                table_lines.append(content_lines[index])
                index += 1
            add_table(document, parse_table(table_lines))
            body_has_content = True
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            markdown_level = len(heading.group(1))
            level = max(1, markdown_level - 1)
            text = heading.group(2).strip()
            normalized = text.casefold()
            if in_abstract and normalized not in {"resumen", "abstract"}:
                document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                document.add_paragraph(title, style="Heading 1")
                body_has_content = True
                in_abstract = False
            if normalized in {"resumen", "abstract"}:
                if body_has_content:
                    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                in_abstract = True
                in_references = False
            elif normalized == "referencias":
                if body_has_content:
                    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                in_references = True
                in_abstract = False
            elif normalized.startswith("apéndice") or normalized.startswith("apendice"):
                if body_has_content:
                    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                in_references = False
                in_abstract = False
            if level <= 3:
                document.add_paragraph(text, style=f"Heading {level}")
            else:
                add_inline_heading(document, text, level)
            body_has_content = True
        elif line.startswith("> "):
            paragraph = document.add_paragraph(style="Block Quote")
            add_inline(paragraph, line[2:].strip())
        elif re.match(r"^[-*]\s+", line):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, re.sub(r"^[-*]\s+", "", line))
        elif re.match(r"^\d+\.\s+", line):
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, re.sub(r"^\d+\.\s+", "", line))
        elif line.strip():
            visual_label = re.fullmatch(r"(Tabla|Figura)\s+\d+[A-Z]?", line.strip(), re.I)
            if visual_label:
                kind = visual_label.group(1).casefold()
                style = "Table Number" if kind == "tabla" else "Figure Number"
                document.add_paragraph(line.strip(), style=style)
                pending_visual_title = kind
            elif pending_visual_title:
                style = "Table Title" if pending_visual_title == "tabla" else "Figure Title"
                paragraph = document.add_paragraph(style=style)
                add_inline(paragraph, line.strip().strip("*"))
                pending_visual_title = None
            elif line.casefold().startswith("nota."):
                paragraph = document.add_paragraph(style="Note")
                label, rest = line[:5], line[5:]
                label_run = paragraph.add_run(label)
                label_run.italic = True
                add_inline(paragraph, rest)
            elif in_abstract and line.casefold().startswith("palabras clave:"):
                paragraph = document.add_paragraph(style="Keywords")
                label, _, rest = line.partition(":")
                label_run = paragraph.add_run(label + ":")
                label_run.italic = True
                add_inline(paragraph, rest)
            else:
                style = "References" if in_references else ("Abstract" if in_abstract else "Normal")
                paragraph = document.add_paragraph(style=style)
                add_inline(paragraph, line.strip())
            body_has_content = True
        index += 1
    document.core_properties.title = title
    document.core_properties.subject = "Trabajo académico en formato APA 7"
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("markdown", type=Path)
    parser.add_argument("output", type=Path, nargs="?")
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--title")
    parser.add_argument("--author")
    parser.add_argument("--institution")
    parser.add_argument("--course")
    parser.add_argument("--instructor")
    parser.add_argument("--due-date", dest="due_date")
    args = parser.parse_args()
    output = args.output or args.markdown.with_name(sanitize_filename(args.title or args.markdown.stem) + ".docx")
    metadata = {key: value for key, value in vars(args).items() if isinstance(value, str) and value}
    convert(args.markdown, output, args.template, metadata)
    print(output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
